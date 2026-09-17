import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.comments import Comment

M = json.load(open('model.json'))
TODAY = '2026-09-17'
FONT = 'Tahoma'

# ---- design tokens -------------------------------------------------
INK, MUTED, LINE = '1A1D23', '6B7280', 'E3E6EA'
CARD, PAGE = 'FFFFFF', 'F4F6F8'
BRAND, BRAND_D = '1F6FEB', '0B3B8C'
GREEN, YELLOW, RED = '16A34A', 'D97706', 'DC2626'
GREEN_BG, YELLOW_BG, RED_BG = 'E7F6EC', 'FDF3E2', 'FCEBEC'
HDR_BG = '111827'

THB = '฿#,##0;[Red]-฿#,##0'
THB2 = '฿#,##0.00;[Red]-฿#,##0.00'
USD = '$#,##0.00;[Red]-$#,##0.00'
USD4 = '$#,##0.0000'
PCT = '0.0%;[Red]-0.0%'
PCT2 = '0.00%'
QTY = '#,##0.000000'
DATE = 'yyyy-mm-dd'
NUM = '#,##0.00'

def F(sz=10, b=False, c=INK, i=False):
    return Font(name=FONT, size=sz, bold=b, color=c, italic=i)
def fill(c): return PatternFill('solid', fgColor=c)
def thin(c=LINE):
    s = Side('thin', color=c); return Border(left=s, right=s, top=s, bottom=s)
def box(c=LINE):
    s = Side('medium', color=c); return Border(left=s, right=s, top=s, bottom=s)

def page(ws, tab=None):
    ws.sheet_view.showGridLines = False
    if tab: ws.sheet_properties.tabColor = tab

def title(ws, row, text, sub='', span=10, col=2):
    c = ws.cell(row, col, text); c.font = F(16, True, INK)
    ws.row_dimensions[row].height = 26
    if sub:
        s = ws.cell(row+1, col, sub); s.font = F(9, False, MUTED)
        ws.row_dimensions[row+1].height = 14

def widths(ws, spec, start=2):
    for i, w in enumerate(spec):
        ws.column_dimensions[get_column_letter(start+i)].width = w

def header(ws, row, cols, start=2):
    for i, h in enumerate(cols):
        c = ws.cell(row, start+i, h)
        c.font = F(9, True, 'FFFFFF'); c.fill = fill(HDR_BG)
        c.alignment = Alignment('center', 'center', wrap_text=True)
        c.border = thin('2A3342')
    ws.row_dimensions[row].height = 30

def mktable(ws, name, row, ncols, nrows, start=2, style='TableStyleLight9'):
    """Register an Excel Table over the header+body range."""
    a = f"{get_column_letter(start)}{row}:{get_column_letter(start+ncols-1)}{row+max(nrows,1)}"
    t = Table(displayName=name, ref=a)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True,
                                      showColumnStripes=False,
                                      showFirstColumn=False, showLastColumn=False)
    ws.add_table(t)
    return t

def put(ws, r, c, v, nf=None, f=None, al=None, bd=True, bg=None):
    cell = ws.cell(r, c, v)
    if nf: cell.number_format = nf
    cell.font = f or F(9)
    if al: cell.alignment = Alignment(horizontal=al, vertical='center')
    else: cell.alignment = Alignment(vertical='center')
    if bd: cell.border = thin()
    if bg: cell.fill = fill(bg)
    return cell

def S(key):
    """XLOOKUP into the settings table."""
    return f'XLOOKUP("{key}",tblSettings[Key],tblSettings[Value])'
