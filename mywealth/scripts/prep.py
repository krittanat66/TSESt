"""Normalize the legacy 'edited.xlsx' workbook into the MY WEALTH data model."""
import openpyxl, json, csv, datetime, collections

SRC = '../edited.xlsx'
wb = openpyxl.load_workbook(SRC, data_only=True)
rows = lambda n: list(wb[n].iter_rows(values_only=True))

def iso(dt):
    d, tm = str(dt).split(' ', 1)
    dd, mm, yy = d.split('/')
    return f"{int(yy)-543:04d}-{int(mm):02d}-{int(dd):02d}", tm

# ---------------------------------------------------------------- FX
fx_conv = []
for r in rows('DIME EX')[4:]:
    if not r[0]: continue
    d, tm = iso(r[0])
    fx_conv.append(dict(date=d, time=tm, thb=float(r[2]), usd=float(r[4]),
                        rate=float(r[5]), ref=r[6] or '', src='DIME EX'))

# ---------------------------------------------------------------- transfers
transfers = []
for r in rows('ประวัติโยกเงิน')[4:]:
    if not r[0]: continue
    d, tm = iso(r[0])
    transfers.append(dict(date=d, time=tm, acct=r[1], kind=r[2],
                          amt=float(r[3]), cur=r[4]))

# two FX conversions present in the transfer ledger but MISSING from DIME EX
known = {(f['date'], round(f['thb'], 2)) for f in fx_conv}
for t in transfers:
    if t['acct'] == 'Dime! Save' and t['kind'] == 'โอนเงิน/ชำระเงิน':
        key = (t['date'], round(-t['amt'], 2))
        if key not in known:
            mate = [x for x in transfers if x['date'] == t['date']
                    and x['acct'] == 'Dime! FCD' and x['kind'] == 'โอนเข้า']
            if mate:
                usd = mate[0]['amt']
                fx_conv.append(dict(date=t['date'], time=t['time'], thb=-t['amt'],
                                    usd=usd, rate=-t['amt']/usd, ref='',
                                    src='DERIVED from ประวัติโยกเงิน (missing in DIME EX)'))
fx_conv.sort(key=lambda x: (x['date'], x['time']))

# ---------------------------------------------------------------- US stock purchases
purchases = []
ws = wb['USA STOCK HIS 2']; rr = list(ws.iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rr) if r[0] and 'วันที่ (ค.ศ.)' in str(r[0]))
for r in rr[hi+1:]:
    if r[0] is None: continue
    d = r[0]
    if isinstance(d, (datetime.date, datetime.datetime)):
        d = d.strftime('%Y-%m-%d')
    else:
        s0 = str(d)
        if 'T' in s0:   # UTC timestamp -> Asia/Bangkok (UTC+7) calendar date
            dt = datetime.datetime.strptime(s0.split('.')[0], '%Y-%m-%dT%H:%M:%S')
            d = (dt + datetime.timedelta(hours=7)).strftime('%Y-%m-%d')
        else:
            d = s0
    purchases.append(dict(date=d, ticker=r[1], usd=float(r[2]),
                          price=float(r[3]), shares=float(r[4]), note=r[5] or ''))
purchases.sort(key=lambda x: (x['date'], x['ticker']))

# ---------------------------------------------------------------- US stock sells
sells = [
    dict(date='2025-08-06', ticker='AMPX', shares=7.03135,  price=7.8334,  usd=55.0794,  avgcost=6.3897,   pl=10.1512),
    dict(date='2025-10-21', ticker='JEPQ', shares=7.10902,  price=57.953,  usd=411.989,  avgcost=55.4677,  pl=17.6680),
    dict(date='2026-03-03', ticker='SCHD', shares=4.0,      price=31.40,   usd=125.60,   avgcost=27.6244,  pl=15.1025),
    dict(date='2026-03-06', ticker='AMPX', shares=1.0,      price=16.2368, usd=16.2368,  avgcost=7.58484,  pl=8.65196),
    dict(date='2026-04-21', ticker='UNH',  shares=0.795472, price=349.69,  usd=278.168,  avgcost=328.696,  pl=16.6999),
]

# ---------------------------------------------------------------- GOLD (rebuilt from Dime screenshots IMG_7097-7102)
gold_raw = [
 ("2025-07-28","22:45:48","BUY", 0.0924,3315.49,"THB",9999.72),("2025-08-08","13:26:15","SELL",0.0924,3393.67,"USD",313.57),
 ("2025-08-11","22:31:01","BUY", 0.0149,3351.10,"USD",49.94),  ("2025-08-12","19:12:46","BUY", 0.0149,3345.16,"USD",49.85),
 ("2025-08-18","09:38:52","BUY", 0.0149,3344.54,"USD",49.84),  ("2025-08-28","18:39:37","SELL",0.0447,3406.51,"USD",152.27),
 ("2025-09-30","15:52:26","BUY", 0.0052,3820.91,"USD",19.87),  ("2025-09-30","16:28:46","BUY", 0.0052,3801.30,"USD",19.77),
 ("2025-10-03","00:24:35","BUY", 0.0159,3847.49,"THB",1999.80),("2025-10-10","08:48:10","BUY", 0.0025,3980.32,"USD",9.96),
 ("2025-10-16","22:07:26","BUY", 0.0023,4268.83,"USD",9.82),   ("2025-10-28","13:48:54","SELL",0.0311,3944.48,"USD",122.67),
 ("2025-10-30","09:09:20","BUY", 0.0203,3935.60,"USD",79.90),  ("2025-11-04","09:51:45","BUY", 0.0125,3978.66,"USD",49.74),
 ("2025-12-01","12:58:58","SELL",0.0328,4227.69,"USD",138.66), ("2026-01-22","08:52:20","BUY", 0.0066,4802.13,"THB",999.77),
 ("2026-01-29","19:29:52","BUY", 0.0090,5541.53,"USD",49.88),  ("2026-02-13","07:19:02","BUY", 0.0050,4943.71,"USD",24.72),
 ("2026-02-23","15:52:08","BUY", 0.0038,5133.73,"USD",19.51),  ("2026-03-05","22:54:48","BUY", 0.0039,5096.97,"USD",19.88),
 ("2026-03-16","09:57:29","SELL",0.0283,5005.73,"USD",141.66), ("2026-03-20","08:35:48","BUY", 0.0064,4675.32,"USD",29.93),
 ("2026-04-01","19:33:39","SELL",0.0064,4740.01,"USD",30.33),  ("2026-05-01","11:49:15","BUY", 0.0064,4611.99,"USD",29.52),
 ("2026-05-28","08:22:31","BUY", 0.0225,4431.33,"USD",99.71),  ("2026-06-08","13:04:23","BUY", 0.0116,4303.98,"USD",49.93),
 ("2026-06-10","23:26:55","BUY", 0.0072,4131.96,"USD",29.76),  ("2026-08-20","08:36:01","SELL",0.0477,4498.43,"USD",214.57),
]
MISSING_IN_LEGACY = {("2026-03-16","SELL"), ("2026-03-20","BUY")}
gold, pos, cost = [], 0.0, 0.0
for d, tm, t, oz, px, cur, amt in gold_raw:
    note = 'Recovered from Dime screenshot - MISSING in legacy GOLD HIS' if (d, t) in MISSING_IN_LEGACY else ''
    if t == 'BUY':
        pos += oz; cost += oz*px; pl = None; cb = oz*px
    else:
        avg = cost/pos if pos else 0.0
        cb = oz*avg; pl = oz*px - cb
        cost -= cb; pos -= oz
        if abs(pos) < 1e-9: pos, cost = 0.0, 0.0
    gold.append(dict(date=d, time=tm, type=t, oz=oz, price=px, cur=cur, amt=amt,
                     costbasis=cb, pl=pl, pos_after=pos, note=note))

# ---------------------------------------------------------------- PVD history
pvd = []
for r in rows('PVD HIS'):
    if not r[0] or not isinstance(r[1], (int, float)): continue
    label = str(r[0])
    pvd.append(dict(label=label, emp=float(r[1]), empben=float(r[2]),
                    er=float(r[3]), erben=float(r[4]), total=float(r[5])))
TH_M = {'มกราคม':1,'กุมภาพันธ์':2,'มีนาคม':3,'เมษายน':4,'พฤษภาคม':5,'มิถุนายน':6,
        'กรกฎาคม':7,'สิงหาคม':8,'กันยายน':9,'ตุลาคม':10,'พฤศจิกายน':11,'ธันวาคม':12}
for p in pvd:
    mn, yr = p['label'].split()
    p['month'] = f"{int(yr)-543:04d}-{TH_M[mn]:02d}-01"
pvd.sort(key=lambda x: x['month'])

# ---------------------------------------------------------------- US portfolio (reported = OFFICIAL per user)
port = []
for r in rows('USA STOCK')[3:]:
    if not r[0] or r[0] in ('หุ้น', 'รวม') or not isinstance(r[1], (int, float)): continue
    if str(r[0]).startswith(('มูลค่า', '%')): continue
    port.append(dict(ticker=r[0], value=float(r[1]), plpct=float(r[2]), pl=float(r[3]),
                     qty=float(r[4]), avgcost=float(r[5]), price=float(r[6])))

out = dict(fx_conv=fx_conv, transfers=transfers, purchases=purchases, sells=sells,
           gold=gold, pvd=pvd, port=port)
json.dump(out, open('model.json', 'w'), ensure_ascii=False, indent=1)
print(f"fx_conv={len(fx_conv)} transfers={len(transfers)} purchases={len(purchases)} "
      f"sells={len(sells)} gold={len(gold)} pvd={len(pvd)} port={len(port)}")
print("gold final position:", gold[-1]['pos_after'],
      " realized total:", round(sum(g['pl'] for g in gold if g['pl'] is not None), 3))
print("fx weighted avg:", round(sum(f['thb'] for f in fx_conv)/sum(f['usd'] for f in fx_conv), 4))
